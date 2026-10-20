"""Write orders into the user's .xlsx template, keeping its header row, column order and formatting."""

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.worksheet import Worksheet

from converter.mapping import flatten_orders
from converter.models import Order

# Matches how the user's sample output displays dates ("November 24, 2025") while keeping a real date cell.
DATE_DISPLAY_FORMAT = "mmmm d, yyyy"


def find_header_sheet(workbook, required_headers: set[str]) -> Worksheet:
    for worksheet in workbook.worksheets:
        first_row = {cell.value for cell in worksheet[1] if cell.value is not None}
        if required_headers <= first_row:
            return worksheet
    raise ValueError(f"Template has no sheet whose first row contains: {sorted(required_headers)}")


def read_template_headers(workbook) -> tuple[str, list[str]]:
    """Header sheet name and its text headers: the active sheet, else the first sheet with a non-empty first row."""
    for worksheet in sorted(workbook.worksheets, key=lambda sheet: sheet is not workbook.active):
        # Non-text header cells cannot be column_map keys (find_header_sheet compares cell values to them).
        headers = [cell.value for cell in worksheet[1] if isinstance(cell.value, str) and cell.value.strip()]
        if headers:
            return worksheet.title, headers
    raise ValueError("Template has no sheet with text headers in its first row")


def generate_blank_template(headers: list[str], sheet_name: str) -> bytes:
    """A one-sheet .xlsx whose first row is `headers`, for templates built in the UI instead of uploaded."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for column_index, header in enumerate(headers, start=1):
        # Force plain text so a header typed as "=..." is not saved as a live formula.
        worksheet.cell(row=1, column=column_index, value=header).data_type = "s"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_xlsx(orders: list[Order], template_path: Path, output_path: Path, column_map: dict[str, str]) -> Path:
    workbook = load_workbook(template_path)
    worksheet = find_header_sheet(workbook, set(column_map))
    # Templates are often prior exports; leftover sheets would carry stale rows into the next export.
    for sheet_name in workbook.sheetnames:
        if sheet_name != worksheet.title:
            workbook.remove(workbook[sheet_name])
    workbook.active = worksheet
    headers = [cell.value for cell in worksheet[1]]
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    for row_index, values in enumerate(flatten_orders(orders, headers, column_map), start=2):
        for column_index, value in enumerate(values, start=1):
            if value is None:
                continue
            if isinstance(value, str):
                # Control characters are invalid in xlsx XML and make openpyxl raise IllegalCharacterError.
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            if isinstance(value, str):
                # openpyxl turns a leading "=" into a live formula; force plain text so OCR/user text never executes.
                cell.data_type = "s"
            if isinstance(value, datetime):
                cell.number_format = DATE_DISPLAY_FORMAT

    workbook.save(output_path)
    return output_path
