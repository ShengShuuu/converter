"""App-level flow through FastAPI's TestClient; OCR stubbed with the synthetic TSV fixtures, jobs run inline."""

import csv
import io
import json
import logging
import re
import sqlite3
import zipfile
from collections import deque
from concurrent.futures import Executor, Future
from contextlib import closing
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path

import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from converter.app import routes_conversions
from converter.app.main import create_app
from converter.extraction import extract_order, parse_tesseract_tsv
from converter.mapping import ODOO_SALE_ORDER_COLUMNS
from converter.ocr.pipeline import PageOcrResult
from converter.ocr.tools import OcrToolError
from converter.profiles import load_profile
from converter.store import repositories as repo
from converter.store.db import INTERRUPTED_ERROR, PAGE_INTERRUPTED_ERROR, connect, init_db
from converter.worker import runner

FIXTURES = Path(__file__).parent / "fixtures"
APP_DIR = Path(__file__).parent.parent / "converter" / "app"
PDF_BYTES = b"%PDF-1.4\n% synthetic stub, OCR is monkeypatched\n"
HX = {"HX-Request": "true"}
GARBAGE_TSV = (
    "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
    "5\t1\t1\t1\t1\t1\t150\t600\t170\t60\t96.5\tCOLLECTION\n"
    "5\t1\t1\t1\t1\t2\t340\t600\t170\t60\t95.0\tRECEIPT\n"
)


class InlineExecutor(Executor):
    """Runs submitted jobs immediately so tests see the final status without threads."""

    def submit(self, fn, /, *args, **kwargs):
        future: Future = Future()
        future.set_result(fn(*args, **kwargs))
        return future


def fixture_words(name: str):
    return parse_tesseract_tsv((FIXTURES / name).read_text(encoding="utf-8"))


@pytest.fixture
def ocr_pages(monkeypatch, tmp_path):
    """Stub OCR returning the S08824 fixture words; set pages / words[n] / errors[n] to shape a batch."""
    raster = tmp_path / "raster.png"
    Image.fromarray(np.full((40, 30, 3), 255, np.uint8)).save(raster)
    default_words = fixture_words("odoo_sale_order_S08824.tsv")
    state = {"pages": 1, "words": {}, "errors": {}, "calls": []}

    def result(page_number):
        state["calls"].append(page_number)
        if page_number in state["errors"]:
            raise state["errors"][page_number]
        words = state["words"].get(page_number, default_words)
        return PageOcrResult(page_number, words, None, 0, raster, raster, {"ocr": 0.1})

    monkeypatch.setattr(runner, "page_count", lambda path: state["pages"])
    monkeypatch.setattr(runner, "ocr_pdf_page", lambda path, page, work: result(page))
    monkeypatch.setattr(runner, "ocr_image_file", lambda path, work: result(1))
    return state


@pytest.fixture
def three_pages(ocr_pages):
    """Page 1: S08824 Sales Order, page 2: another document, page 3: the S07002 letterhead Sales Order."""
    ocr_pages["pages"] = 3
    ocr_pages["words"] = {2: parse_tesseract_tsv(GARBAGE_TSV), 3: fixture_words("odoo_sale_order_S07002_letterhead.tsv")}
    return ocr_pages


@pytest.fixture
def data_dir(tmp_path):
    return tmp_path / "data"


@pytest.fixture
def client(data_dir, ocr_pages):
    with TestClient(create_app(data_dir, executor_factory=InlineExecutor)) as test_client:
        yield test_client


def db_of(data_dir: Path) -> Path:
    return data_dir.resolve() / "converter.db"


def page_of(data_dir: Path, conversion_id: int, page_number: int = 1) -> repo.Page:
    return repo.get_page(db_of(data_dir), conversion_id, page_number)


def order_url(conversion_id: int, page_number: int = 1) -> str:
    return f"/conversions/{conversion_id}/pages/{page_number}"


def upload_pdf(client, name="order.pdf", content=PDF_BYTES):
    return client.post("/conversions", files={"file": (name, content, "application/pdf")}, headers=HX)


def uploaded_conversion(client, name="order.pdf") -> int:
    response = upload_pdf(client, name)
    assert response.status_code == 200
    return int(response.headers["HX-Redirect"].rsplit("/", 1)[1])


def upload_template(client, name="Odoo export", next_url=""):
    content = (FIXTURES / "sample_template.xlsx").read_bytes()
    return client.post(
        "/templates",
        data={"name": name, "next": next_url},
        files={"file": ("template.xlsx", content, "application/octet-stream")},
        follow_redirects=False,
    )


def choose_uploaded_template(client, conversion_id: int) -> int:
    template_id = int(upload_template(client).headers["location"].split("?")[0].rsplit("/", 1)[1])
    response = client.post(f"/conversions/{conversion_id}/template", data={"template_id": str(template_id)},
                           follow_redirects=False)
    assert response.headers["location"] == f"/conversions/{conversion_id}/preview"
    return template_id


def ready_for_export(client, data_dir, name="order.pdf") -> int:
    conversion_id = uploaded_conversion(client, name)
    choose_uploaded_template(client, conversion_id)
    return conversion_id


def confirm_everything(client, conversion_id: int, page_number: int = 1) -> None:
    review = client.get(order_url(conversion_id, page_number)).text
    for key in set(re.findall(r'/confirm/([\w.]+)"', review)):
        assert client.post(f"{order_url(conversion_id, page_number)}/confirm/{key}", headers=HX).status_code == 200


def test_full_flow_upload_review_template_preview(client, data_dir):
    assert client.get("/").status_code == 200
    conversion_id = uploaded_conversion(client)
    db_path = db_of(data_dir)
    assert repo.get_conversion(db_path, conversion_id).status == "done"
    assert page_of(data_dir, conversion_id).status == "done"

    review = client.get(order_url(conversion_id))
    assert review.status_code == 200 and "S08824" in review.text
    assert 'id="review-summary"' in review.text and 'id="totals-check"' in review.text
    assert client.get(f"{order_url(conversion_id)}/page-image").headers["content-type"] == "image/png"

    extracted_total = page_of(data_dir, conversion_id).extracted_order.total
    patched = client.patch(f"{order_url(conversion_id)}/order/total", data={"value": "9,999.00"}, headers=HX)
    assert patched.status_code == 200
    assert 'id="cell-total"' in patched.text and 'value="9999.00"' in patched.text
    assert 'id="review-summary" hx-swap-oob="true"' in patched.text and 'id="totals-check" hx-swap-oob="true"' in patched.text
    stored = page_of(data_dir, conversion_id)
    assert stored.order.total == Decimal("9999.00") and "total" in stored.confirmed_fields
    assert stored.extracted_order.total == extracted_total != Decimal("9999.00")

    confirmed = client.post(f"{order_url(conversion_id)}/confirm/order_date", headers=HX)
    assert confirmed.status_code == 200 and "tier--checked" in confirmed.text
    assert "order_date" in page_of(data_dir, conversion_id).confirmed_fields

    line_edit = client.patch(f"{order_url(conversion_id)}/lines/1/subtotal", data={"value": "5760"}, headers=HX)
    assert line_edit.status_code == 200 and 'id="cell-line-1-subtotal"' in line_edit.text

    template_response = upload_template(client, next_url=f"/conversions/{conversion_id}/template")
    assert template_response.status_code == 303
    mapping_url = template_response.headers["location"]
    assert mapping_url.startswith("/templates/") and "next=/conversions/" in mapping_url
    mapping_page = client.get(mapping_url)
    assert mapping_page.status_code == 200
    assert '<option value="line.quantity" selected>' in mapping_page.text

    template_id = int(mapping_url.split("?")[0].rsplit("/", 1)[1])
    form = {f"map_{index}": path for index, path in enumerate(ODOO_SALE_ORDER_COLUMNS.values())}
    form["map_8"] = ""  # leave Order Lines/Taxes blank
    form["next"] = f"/conversions/{conversion_id}/template"
    saved = client.post(f"/templates/{template_id}/mapping", data=form, follow_redirects=False)
    assert saved.headers["location"] == f"/conversions/{conversion_id}/template"
    assert "Order Lines/Taxes" not in repo.get_template(db_path, template_id).column_map

    chosen = client.post(f"/conversions/{conversion_id}/template", data={"template_id": str(template_id)},
                         follow_redirects=False)
    assert chosen.headers["location"] == f"/conversions/{conversion_id}/preview"
    preview = client.get(f"/conversions/{conversion_id}/preview")
    assert preview.status_code == 200
    assert "PX-S100 / 100W PX-S SOLAR FLOOD LIGHT" in preview.text and "Output Tax (12%)" not in preview.text


def test_review_context_carries_order_url_and_navigation(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)

    first = client.get(order_url(conversion_id, 1)).context
    third = client.get(order_url(conversion_id, 3)).context

    assert first["order_url"] == order_url(conversion_id, 1) and first["page"].page_number == 1
    assert (first["done_page_position"], first["done_page_count"]) == (1, 2)
    assert first["prev_page_url"] is None and first["next_page_url"] == order_url(conversion_id, 3)
    assert third["prev_page_url"] == order_url(conversion_id, 1) and third["next_page_url"] is None
    assert third["batch_url"] == f"/conversions/{conversion_id}"
    assert third["order_cells"]["order_ref"].value == "S07002"


class _RequestCollector(HTMLParser):
    """Every request a rendered element would make: (method, url, form data it sends)."""

    URL_ATTRIBUTES = {"hx-get": "GET", "hx-post": "POST", "hx-patch": "PATCH", "hx-put": "PUT",
                      "hx-delete": "DELETE", "src": "GET"}

    def __init__(self):
        super().__init__()
        self.requests: list[tuple[str, str, dict]] = []

    def handle_starttag(self, tag, attrs):
        attributes = {name: value or "" for name, value in attrs}
        for attribute, method in self.URL_ATTRIBUTES.items():
            if attribute not in attributes:
                continue
            if "hx-vals" in attributes:
                data = json.loads(attributes["hx-vals"])
            elif method == "PATCH":
                # A field input sends its own value on change.
                data = {"value": attributes.get("value", "")}
            else:
                data = {}
            self.requests.append((method, attributes[attribute], data))


def rendered_requests(html: str) -> list[tuple[str, str, dict]]:
    collector = _RequestCollector()
    collector.feed(html)
    return collector.requests


def test_every_rendered_review_url_is_page_scoped_and_resolves(client, data_dir, three_pages):
    """Renders templates (unlike the direct route calls above), so a hardcoded conversion-level URL or a missing
    order_url in some handler's template context fails here even when every route works on its own."""
    conversion_id = uploaded_conversion(client)
    # Page 3, not 1: a URL built from the wrong page number must not pass by coincidence.
    prefix = order_url(conversion_id, 3)
    # Force a lines/total mismatch so "Use sum of lines" renders too (LOW-A: it's otherwise never seen).
    client.patch(f"{prefix}/order/total", data={"value": "1.00"}, headers=HX)
    review = client.get(prefix)
    assert review.status_code == 200
    seen: set[tuple[str, str]] = set()
    pending: deque[tuple[str, str, dict]] = deque()
    deletes: list[tuple[str, str, dict]] = []

    def collect(html: str, source: str) -> None:
        for method, url, data in rendered_requests(html):
            if url.startswith("/static/"):
                continue
            assert url.startswith(f"{prefix}/"), f"{source} rendered {method} {url!r}, not under {prefix}/"
            if (method, url) not in seen:
                seen.add((method, url))
                (deletes if method == "DELETE" else pending).append((method, url, data))

    collect(review.text, f"GET {prefix}")
    while pending or deletes:
        if pending:
            method, url, data = pending.popleft()
        else:
            # Highest line index first, so each delete still targets an existing line.
            deletes.sort(key=lambda request: int(request[1].rsplit("/", 1)[1]))
            method, url, data = deletes.pop()
        response = client.request(method, url, data=data or None, headers=HX)
        assert response.status_code == 200, f"replaying {method} {url} ({data}) returned {response.status_code}"
        if response.headers["content-type"].startswith("text/html"):
            collect(response.text, f"{method} {url}")

    replayed = {(method, url.removeprefix(prefix)) for method, url in seen}
    assert ("GET", "/page-image") in replayed and ("POST", "/lines") in replayed
    assert {"PATCH /order", "PATCH /lines", "POST /confirm", "DELETE /lines"} <= {
        f"{method} /{path.split('/')[1]}" for method, path in replayed
    }


def test_invalid_edit_is_not_saved_and_keeps_typed_value(client, data_dir):
    conversion_id = uploaded_conversion(client)
    before = page_of(data_dir, conversion_id).order

    response = client.patch(f"{order_url(conversion_id)}/order/total", data={"value": "12,34.5"}, headers=HX)

    assert response.status_code == 422
    assert 'value="12,34.5"' in response.text and "Enter a number" in response.text
    stored = page_of(data_dir, conversion_id)
    assert stored.order == before and "total" not in stored.confirmed_fields


def test_add_and_delete_lines_remap_confirmations(client, data_dir):
    conversion_id = uploaded_conversion(client)
    url = order_url(conversion_id)
    client.post(f"{url}/confirm/line.1.subtotal", headers=HX)

    added = client.post(f"{url}/lines", headers=HX)
    assert added.status_code == 200 and 'id="line-items"' in added.text
    assert len(page_of(data_dir, conversion_id).order.lines) == 3

    deleted = client.delete(f"{url}/lines/0", headers=HX)
    assert deleted.status_code == 200
    stored = page_of(data_dir, conversion_id)
    assert len(stored.order.lines) == 2
    assert stored.order.lines[0].product == "PX-S100 / 100W PX-S SOLAR FLOOD LIGHT"
    assert stored.confirmed_fields == {"line.0.subtotal"}

    assert client.delete(f"{url}/lines/5", headers=HX).status_code == 409
    assert client.patch(f"{url}/lines/9/product", data={"value": "x"}).status_code == 409
    assert client.patch(f"{url}/order/colour", data={"value": "x"}).status_code == 404
    assert client.post(f"{url}/confirm/line.x.total").status_code == 404


def test_edits_on_one_page_do_not_touch_another(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)
    first_before = page_of(data_dir, conversion_id, 1)

    assert client.patch(f"{order_url(conversion_id, 3)}/order/total", data={"value": "1.00"}, headers=HX).status_code == 200
    assert client.post(f"{order_url(conversion_id, 3)}/lines", headers=HX).status_code == 200

    first, third = page_of(data_dir, conversion_id, 1), page_of(data_dir, conversion_id, 3)
    assert first.order == first_before.order and first.confirmed_fields == frozenset()
    assert third.order.total == Decimal("1.00") and third.confirmed_fields == {"total"}


def test_page_routes_check_the_page_exists_and_is_reviewable(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)

    not_order = client.get(order_url(conversion_id, 2), follow_redirects=False)
    assert not_order.status_code == 303 and not_order.headers["location"] == f"/conversions/{conversion_id}"
    assert client.patch(f"{order_url(conversion_id, 2)}/order/total", data={"value": "1"}).status_code == 409
    assert client.post(f"{order_url(conversion_id, 2)}/lines").status_code == 409
    assert client.get(order_url(conversion_id, 9)).status_code == 404
    assert client.patch(f"{order_url(conversion_id, 9)}/order/total", data={"value": "1"}).status_code == 404
    assert client.get(f"{order_url(conversion_id, 9)}/page-image").status_code == 404
    assert client.get(order_url(999, 1)).status_code == 404
    assert client.get(f"{order_url(conversion_id, 2)}/page-image").status_code == 200


def test_export_is_blocked_until_unchecked_fields_are_acknowledged(client, data_dir):
    conversion_id = ready_for_export(client, data_dir)

    blocked = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx"})

    assert blocked.status_code == 422
    assert "have not been checked" in blocked.text
    assert "spreadsheetml" not in blocked.headers["content-type"]


def test_exported_xlsx_is_typed_and_has_one_sheet(client, data_dir, tmp_path):
    conversion_id = ready_for_export(client, data_dir)

    response = client.post(f"/conversions/{conversion_id}/export",
                           data={"format": "xlsx", "acknowledge_unchecked": "1"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="S08824.xlsx"' in response.headers["content-disposition"]
    output = tmp_path / "download.xlsx"
    output.write_bytes(response.content)
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Converted"]
    sheet = workbook["Converted"]
    assert sheet["A2"].value == "S08824"
    assert sheet["B2"].is_date
    assert sheet["H3"].data_type == "n" and sheet["H3"].value == 5760
    assert sheet.max_row == 3


def test_exported_csv_has_bom_and_plain_values(client, data_dir):
    conversion_id = ready_for_export(client, data_dir)

    response = client.post(f"/conversions/{conversion_id}/export",
                           data={"format": "csv", "acknowledge_unchecked": "1"})

    assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
    assert response.content.startswith(b"\xef\xbb\xbf")
    header, first, _ = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert header == list(ODOO_SALE_ORDER_COLUMNS)
    assert first[0] == "S08824" and first[1] == "2026-04-22 10:41:30"


def test_peso_sign_is_display_only(client, data_dir):
    conversion_id = ready_for_export(client, data_dir)
    order = page_of(data_dir, conversion_id).order
    first_subtotal = order.lines[0].subtotal
    assert order.total == Decimal("8360.00") and first_subtotal is not None

    review = client.get(order_url(conversion_id)).text
    assert 'value="8360.00"' in review and "₱" not in "".join(re.findall(r'value="[^"]*"', review))
    assert "<dd>₱8,360.00</dd>" in review and 'aria-describedby="cell-total-currency cell-total-status"' in review
    assert "Total: ₱8,360.00" in client.get(f"/conversions/{conversion_id}").text
    assert f"₱{first_subtotal:,f}" in client.get(f"/conversions/{conversion_id}/preview").text

    exported = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv", "acknowledge_unchecked": "1"})
    assert "₱" not in exported.content.decode("utf-8-sig") and format(first_subtotal, "f") in exported.text


def test_export_proceeds_without_acknowledgement_once_everything_is_checked(client, data_dir):
    conversion_id = ready_for_export(client, data_dir)
    confirm_everything(client, conversion_id)

    response = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv"})

    assert response.status_code == 200 and response.content.startswith(b"\xef\xbb\xbf")


@pytest.mark.parametrize(
    ("name", "content", "message"),
    [
        ("order.exe", PDF_BYTES, "Upload a PDF, PNG, JPG or TIFF file."),
        ("order.pdf", b"\x89PNG\r\n\x1a\n rest", "not a valid PDF"),
        ("photo.jpg", PDF_BYTES, "not a valid JPG"),
        ("order.pdf", b"", "not a valid PDF"),
    ],
)
def test_upload_rejects_wrong_type(client, data_dir, name, content, message):
    response = upload_pdf(client, name, content)
    assert response.status_code == 422 and message in response.text
    assert repo.list_conversions(db_of(data_dir)) == []


def test_upload_rejects_oversized_file(client, monkeypatch):
    monkeypatch.setattr(routes_conversions, "MAX_UPLOAD_BYTES", 10)
    response = upload_pdf(client)
    assert response.status_code == 422 and "larger than 25 MB" in response.text


# def test_multi_page_pdf_shows_failure_with_retry(...)  -- v1 rejected multi-page PDFs; see the batch tests below
# def test_status_polls_while_processing_and_redirects_when_done(...)  -- v1 sent HX-Redirect to review on done


def test_too_many_pages_fails_the_conversion_and_retry_restarts_it(client, data_dir, ocr_pages):
    ocr_pages["pages"] = 51
    conversion_id = uploaded_conversion(client)

    page = client.get(f"/conversions/{conversion_id}")

    assert page.status_code == 200
    assert page.context["conversion"].status == "failed" and "max 50" in page.context["conversion"].error
    assert page.context["pages"] == [] and not page.context["is_processing"] and not page.context["can_export"]

    ocr_pages["pages"] = 2
    retried = client.post(f"/conversions/{conversion_id}/retry", headers=HX)
    assert retried.status_code == 200 and "HX-Redirect" not in retried.headers
    assert repo.get_conversion(db_of(data_dir), conversion_id).status == "done"
    assert [p.status for p in repo.list_pages(db_of(data_dir), conversion_id)] == ["done", "done"]


def test_batch_list_shows_every_page_with_its_outcome(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)

    page = client.get(f"/conversions/{conversion_id}")

    assert page.status_code == 200 and page.template.name == "batch.html"
    context = page.context
    rows = context["pages"]
    assert [(r.page_number, r.status, r.order_ref) for r in rows] == [
        (1, "done", "S08824"), (2, "not_sales_order", None), (3, "done", "S07002"),
    ]
    assert rows[0].review_url == order_url(conversion_id, 1) and rows[1].review_url is None
    assert rows[0].customer and rows[0].line_count == 2 and rows[0].total is not None
    assert rows[0].unchecked_count > 0 and rows[0].needs_attention
    assert rows[1].classifier_score == 0.0 and context["classifier_threshold"] == runner.CLASSIFIER_MIN_RATIO
    assert (rows[0].can_skip, rows[0].can_extract_anyway, rows[0].can_retry) == (True, False, False)
    assert (rows[1].can_skip, rows[1].can_extract_anyway, rows[1].can_retry) == (False, True, False)
    progress = context["progress"]
    assert (progress.total, progress.finished, progress.pending) == (3, 3, 0)
    assert progress.counts["done"] == 2 and progress.counts["not_sales_order"] == 1
    assert context["can_export"] and not context["is_processing"]
    assert context["batch_url"] == f"/conversions/{conversion_id}"
    assert context["template_url"] == f"/conversions/{conversion_id}/template"
    assert "S08824" in page.text and "S07002" in page.text


def test_single_page_upload_also_lands_on_the_batch_list(client, data_dir):
    conversion_id = uploaded_conversion(client)
    page = client.get(f"/conversions/{conversion_id}", follow_redirects=False)
    assert page.status_code == 200 and page.template.name == "batch.html"
    assert [r.status for r in page.context["pages"]] == ["done"]


def test_status_polls_while_processing_and_never_redirects(client, data_dir):
    db_path = db_of(data_dir)
    conversion = repo.create_conversion(db_path, "scan.pdf", ".pdf")
    repo.create_pages(db_path, conversion.id, 2)
    repo.set_status(db_path, conversion.id, "processing")
    repo.set_page_status(db_path, conversion.id, 1, "not_sales_order")

    polling = client.get(f"/conversions/{conversion.id}/status", headers=HX)
    assert polling.status_code == 200 and "HX-Redirect" not in polling.headers
    assert 'hx-trigger="every 2s"' in polling.text

    done_id = uploaded_conversion(client)
    done = client.get(f"/conversions/{done_id}/status", headers=HX)
    assert "HX-Redirect" not in done.headers and 'hx-trigger="every 2s"' not in done.text


def test_skip_page(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)

    skipped = client.post(f"{order_url(conversion_id, 1)}/skip", headers=HX)
    assert skipped.status_code == 200 and "HX-Redirect" not in skipped.headers
    assert page_of(data_dir, conversion_id, 1).status == "skipped"
    assert client.post(f"{order_url(conversion_id, 1)}/skip", headers=HX).status_code == 200

    not_htmx = client.post(f"{order_url(conversion_id, 3)}/skip", follow_redirects=False)
    assert not_htmx.status_code == 303 and not_htmx.headers["location"] == f"/conversions/{conversion_id}"
    assert client.get(order_url(conversion_id, 1), follow_redirects=False).status_code == 303
    # assert client.post(f"{order_url(conversion_id, 2)}/skip", headers=HX).status_code == 409
    not_order = client.post(f"{order_url(conversion_id, 2)}/skip", headers=HX)
    assert not_order.status_code == 409 and "not a Sales Order" in not_order.json()["detail"]
    assert client.post(f"{order_url(conversion_id, 9)}/skip", headers=HX).status_code == 404

    repo.set_page_status(db_of(data_dir), conversion_id, 2, "processing")
    # assert client.post(f"{order_url(conversion_id, 2)}/skip", headers=HX).status_code == 409
    busy = client.post(f"{order_url(conversion_id, 2)}/skip", headers=HX)
    assert busy.status_code == 409 and "still being read" in busy.json()["detail"]
    assert page_of(data_dir, conversion_id, 2).status == "processing"


def test_extract_anyway_reads_saved_words(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)
    three_pages["errors"] = {2: AssertionError("OCR must not run again")}

    response = client.post(f"{order_url(conversion_id, 2)}/extract-anyway", headers=HX)

    assert response.headers["HX-Redirect"] == order_url(conversion_id, 2)
    stored = page_of(data_dir, conversion_id, 2)
    assert stored.status == "done" and stored.order.source_pages == [2]


def test_extract_anyway_undoes_a_skip_and_ignores_other_states(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)
    client.post(f"{order_url(conversion_id, 1)}/skip", headers=HX)

    redone = client.post(f"{order_url(conversion_id, 1)}/extract-anyway", follow_redirects=False)
    assert redone.status_code == 303 and redone.headers["location"] == order_url(conversion_id, 1)
    assert page_of(data_dir, conversion_id, 1).status == "done"

    before = page_of(data_dir, conversion_id, 3)
    ignored = client.post(f"{order_url(conversion_id, 3)}/extract-anyway", headers=HX)
    assert ignored.status_code == 200 and "HX-Redirect" not in ignored.headers
    assert page_of(data_dir, conversion_id, 3).updated_at == before.updated_at


def test_undo_skip_restores_the_reviewed_order_without_re_extracting(client, data_dir, three_pages, monkeypatch):
    conversion_id = uploaded_conversion(client)
    url = order_url(conversion_id, 1)
    assert client.patch(f"{url}/order/customer", data={"value": "MY EDIT"}, headers=HX).status_code == 200
    assert client.post(f"{url}/confirm/order_date", headers=HX).status_code == 200
    assert client.delete(f"{url}/lines/0", headers=HX).status_code == 200
    reviewed = page_of(data_dir, conversion_id, 1)
    assert reviewed.confirmed_fields == {"customer", "order_date"} and len(reviewed.order.lines) == 1

    def no_re_extraction(*args):
        raise AssertionError("undoing a skip must not re-extract the page")

    monkeypatch.setattr(routes_conversions, "extract_from_saved_words", no_re_extraction)
    assert client.post(f"{url}/skip", headers=HX).status_code == 200
    response = client.post(f"{url}/extract-anyway", headers=HX)

    assert response.headers["HX-Redirect"] == url
    restored = page_of(data_dir, conversion_id, 1)
    assert (restored.status, restored.error) == ("done", None)
    assert restored.order == reviewed.order and restored.order.customer == "MY EDIT"
    assert restored.confirmed_fields == reviewed.confirmed_fields
    assert restored.extracted_order == reviewed.extracted_order


def test_extract_anyway_still_re_extracts_pages_without_an_order(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)

    not_order = client.post(f"{order_url(conversion_id, 2)}/extract-anyway", headers=HX)
    assert not_order.headers["HX-Redirect"] == order_url(conversion_id, 2)
    extracted = page_of(data_dir, conversion_id, 2)
    assert extracted.status == "done" and extracted.order is not None and extracted.order == extracted.extracted_order
    assert extracted.confirmed_fields == frozenset()

    # A skipped page that never had an order (e.g. a failed page the user skipped) has nothing to restore.
    db_path = db_of(data_dir)
    with closing(connect(db_path)) as connection:
        connection.execute("UPDATE pages SET status = 'skipped', order_json = NULL, extracted_order_json = NULL"
                           " WHERE conversion_id = ? AND page_number = 2", (conversion_id,))
    skipped_without_order = client.post(f"{order_url(conversion_id, 2)}/extract-anyway", headers=HX)
    assert skipped_without_order.headers["HX-Redirect"] == order_url(conversion_id, 2)
    re_extracted = page_of(data_dir, conversion_id, 2)
    assert re_extracted.status == "done" and re_extracted.order == extracted.order


def test_extract_anyway_without_saved_words_fails_the_page(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)
    runner.saved_words_path(data_dir.resolve(), conversion_id, 2).unlink()

    response = client.post(f"{order_url(conversion_id, 2)}/extract-anyway", headers=HX)

    assert response.status_code == 200 and "HX-Redirect" not in response.headers
    stored = page_of(data_dir, conversion_id, 2)
    assert (stored.status, stored.error) == ("failed", routes_conversions.MISSING_WORDS_ERROR)


def test_retry_one_page_rereads_only_that_page(client, data_dir, three_pages):
    three_pages["errors"] = {2: OcrToolError("pdftoppm exited 99"), 3: OcrToolError("tesseract exited 1")}
    conversion_id = uploaded_conversion(client)
    client.patch(f"{order_url(conversion_id, 1)}/order/customer", data={"value": "EDITED"}, headers=HX)
    assert [p.status for p in repo.list_pages(db_of(data_dir), conversion_id)] == ["done", "failed", "failed"]
    rows = client.get(f"/conversions/{conversion_id}").context["pages"]
    assert rows[1].can_retry and rows[1].error == "pdftoppm exited 99"
    three_pages["errors"] = {}
    three_pages["calls"].clear()

    response = client.post(f"{order_url(conversion_id, 2)}/retry", headers=HX)

    assert response.status_code == 200
    assert three_pages["calls"] == [2]
    assert [p.status for p in repo.list_pages(db_of(data_dir), conversion_id)] == ["done", "not_sales_order", "failed"]
    first = page_of(data_dir, conversion_id, 1)
    assert first.order.customer == "EDITED" and first.confirmed_fields == {"customer"}
    assert repo.get_conversion(db_of(data_dir), conversion_id).status == "done"

    three_pages["calls"].clear()
    assert client.post(f"/conversions/{conversion_id}/retry", headers=HX).status_code == 200
    assert three_pages["calls"] == [3]
    assert page_of(data_dir, conversion_id, 3).status == "done"

    three_pages["calls"].clear()
    client.post(f"{order_url(conversion_id, 1)}/retry", headers=HX)
    client.post(f"/conversions/{conversion_id}/retry", headers=HX)
    assert three_pages["calls"] == []


def test_delete_conversion_removes_files_and_pages(client, data_dir, three_pages):
    conversion_id = uploaded_conversion(client)
    directory = repo.conversion_dir(data_dir.resolve(), conversion_id)
    assert (directory / "source.pdf").is_file()
    assert runner.display_image_path(data_dir.resolve(), conversion_id, 3).is_file()

    response = client.delete(f"/conversions/{conversion_id}", headers=HX)

    assert response.status_code == 200 and response.text == ""
    assert not directory.exists()
    assert client.get(f"/conversions/{conversion_id}").status_code == 404
    assert client.get(order_url(conversion_id, 1)).status_code == 404
    assert repo.list_pages(db_of(data_dir), conversion_id) == []


def test_template_upload_validation(client):
    assert upload_template(client).status_code == 303
    duplicate = upload_template(client)
    assert duplicate.status_code == 422 and "already exists" in duplicate.text
    not_xlsx = client.post("/templates", data={"name": "x"}, files={"file": ("t.xlsx", b"not a zip", "x")})
    assert not_xlsx.status_code == 422 and "not a valid .xlsx" in not_xlsx.text
    corrupt = client.post("/templates", data={"name": "y"}, files={"file": ("t.xlsx", b"PK\x03\x04junk", "x")})
    assert corrupt.status_code == 422 and "could not be read" in corrupt.text


def test_manual_template_creates_row_and_blank_workbook(client, data_dir):
    page = client.get("/templates/new-manual")
    assert page.status_code == 200 and 'action="/templates/manual"' in page.text
    assert 'href="/templates/new-manual"' in client.get("/templates").text

    form = {
        "name": "Built by hand",
        "header": ["Customer", "Notes", "=Total", ""],
        "field": ["order.customer", "", "order.total", ""],
        "next": "/conversions/5",
    }
    response = client.post("/templates/manual", data=form, follow_redirects=False)
    assert response.status_code == 303 and response.headers["location"] == "/conversions/5"

    (template,) = repo.list_templates(db_of(data_dir))
    assert template.name == "Built by hand"
    assert template.headers == ["Customer", "Notes", "=Total"]
    assert template.column_map == {"Customer": "order.customer", "=Total": "order.total"}
    workbook = load_workbook(data_dir.resolve() / template.file_path)
    assert workbook.sheetnames == [template.sheet_name]
    assert [cell.value for cell in workbook[template.sheet_name][1]] == template.headers


@pytest.mark.parametrize(
    ("form", "message"),
    [
        ({"name": "", "header": ["A"], "field": [""]}, "Enter a name"),
        ({"name": "t", "header": ["", " "], "field": ["", ""]}, "at least one column header"),
        ({"name": "t", "header": ["A", "A"], "field": ["", ""]}, "used more than once"),
        ({"name": "t", "header": ["A", ""], "field": ["", "order.customer"]}, "Enter a column header"),
        ({"name": "t", "header": ["A"], "field": ["order.nope"]}, "Unknown field"),
    ],
)
def test_manual_template_validation(client, data_dir, form, message):
    response = client.post("/templates/manual", data=form)
    assert response.status_code == 422 and message in response.text
    assert repo.list_templates(db_of(data_dir)) == []


def test_manual_template_duplicate_name_keeps_typed_columns(client):
    assert upload_template(client, name="Taken").status_code == 303
    response = client.post("/templates/manual", data={"name": "Taken", "header": ["Kept header"], "field": [""]})
    assert response.status_code == 422 and "already exists" in response.text
    assert 'value="Kept header"' in response.text


def test_next_parameter_cannot_redirect_off_site(client):
    template_id = int(upload_template(client, next_url="//evil.example").headers["location"].rsplit("/", 1)[1])
    response = client.post(f"/templates/{template_id}/mapping", data={"next": "https://evil.example"},
                           follow_redirects=False)
    assert response.headers["location"] == "/templates"


@pytest.mark.parametrize("next_url", ["//evil.example", "/\t/evil.example", "/\n/evil.example", "/\\evil.example",
                                      "/x\\..\\/evil.example", "https://evil.example"])
def test_next_values_that_browsers_turn_off_site_are_dropped(client, next_url):
    template_id = int(upload_template(client).headers["location"].rsplit("/", 1)[1])
    page = client.get(f"/templates/{template_id}", params={"next": next_url})
    assert "evil.example" not in page.text
    saved = client.post(f"/templates/{template_id}/mapping", data={"next": next_url}, follow_redirects=False)
    assert saved.headers["location"] == "/templates"


def test_relative_next_path_is_kept(client):
    template_id = int(upload_template(client).headers["location"].rsplit("/", 1)[1])
    assert 'href="/conversions/5"' in client.get(f"/templates/{template_id}", params={"next": "/conversions/5"}).text
    saved = client.post(f"/templates/{template_id}/mapping", data={"next": "/conversions/5"}, follow_redirects=False)
    assert saved.headers["location"] == "/conversions/5"


def test_pasted_control_character_does_not_break_xlsx_export(client, data_dir, tmp_path):
    conversion_id = ready_for_export(client, data_dir)
    edited = client.patch(f"{order_url(conversion_id)}/order/customer", data={"value": "ACME\x0bCORP"}, headers=HX)
    assert edited.status_code == 200

    response = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx", "acknowledge_unchecked": "1"})

    assert response.status_code == 200
    output = tmp_path / "download.xlsx"
    output.write_bytes(response.content)
    assert load_workbook(output)["Converted"]["C2"].value == "ACMECORP"


def test_manual_template_exports_through_the_unchanged_xlsx_path(client, data_dir, tmp_path):
    conversion_id = uploaded_conversion(client)
    form = {"name": "Manual", "header": ["Product", "Note"], "field": ["line.product", ""]}
    assert client.post("/templates/manual", data=form, follow_redirects=False).status_code == 303
    (template,) = repo.list_templates(db_of(data_dir))
    chosen = client.post(f"/conversions/{conversion_id}/template", data={"template_id": str(template.id)},
                         follow_redirects=False)
    assert chosen.status_code == 303

    response = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx", "acknowledge_unchecked": "1"})

    assert response.status_code == 200
    output = tmp_path / "download.xlsx"
    output.write_bytes(response.content)
    sheet = load_workbook(output)[template.sheet_name]
    assert [cell.value for cell in sheet[1]] == ["Product", "Note"]
    assert sheet["A2"].value and sheet["B2"].value is None


def test_manual_template_custom_column_writes_its_text_on_every_row(client, data_dir, tmp_path):
    conversion_id = uploaded_conversion(client)
    form = {
        "name": "Manual custom",
        "header": ["Product", "Status", "Note"],
        "field": ["line.product", "__custom__", ""],
        # Text on a non-custom row is ignored; the custom text is trimmed.
        "field_text": ["", "  N/A pending\x0b ", "ignored"],
    }
    assert client.post("/templates/manual", data=form, follow_redirects=False).status_code == 303
    (template,) = repo.list_templates(db_of(data_dir))
    assert template.column_map == {"Product": "line.product", "Status": "literal:N/A pending"}
    client.post(f"/conversions/{conversion_id}/template", data={"template_id": str(template.id)})

    preview = client.get(f"/conversions/{conversion_id}/preview")
    assert preview.status_code == 200 and "N/A pending" in preview.text

    response = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx", "acknowledge_unchecked": "1"})
    assert response.status_code == 200
    output = tmp_path / "download.xlsx"
    output.write_bytes(response.content)
    sheet = load_workbook(output)[template.sheet_name]
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    line_count = len(page_of(data_dir, conversion_id).order.lines)
    assert line_count > 1 and len(data_rows) == line_count
    assert all(product for product, _, _ in data_rows)
    assert [status for _, status, _ in data_rows] == ["N/A pending"] * line_count
    assert [note for _, _, note in data_rows] == [None] * line_count
    # A custom value is plain text in the sheet, never a formula, regardless of what it looks like.
    assert sheet["B2"].data_type == "s"
    with zipfile.ZipFile(output) as workbook_zip:
        sheet_xml = next(name for name in workbook_zip.namelist() if name.startswith("xl/worksheets/sheet"))
        assert "<f>" not in workbook_zip.read(sheet_xml).decode("utf-8")

    csv_response = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv", "acknowledge_unchecked": "1"})
    csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert [row[1] for row in csv_rows[1:]] == ["N/A pending"] * line_count
    assert [row[2] for row in csv_rows[1:]] == [""] * line_count


@pytest.mark.parametrize(
    ("form", "message"),
    [
        (
            {"name": "t", "header": ["Status"], "field": ["__custom__"], "field_text": ["  "]},
            "Enter a value for the custom column &#34;Status&#34;, or choose a different option.",
        ),
        (
            {"name": "t", "header": ["Status"], "field": ["__custom__"], "field_text": ["x" * 256]},
            "Custom values can be at most 255 characters.",
        ),
        (
            {"name": "t", "header": ["", "B"], "field": ["__custom__", ""], "field_text": ["Pending", ""]},
            "Enter a column header for the column filled with &#34;Pending&#34;.",
        ),
        (
            {"name": "t", "header": ["A", "B"], "field": ["", ""], "field_text": [""]},
            "Every column needs a header and a field choice.",
        ),
    ],
)
def test_manual_template_custom_column_validation(client, data_dir, form, message):
    response = client.post("/templates/manual", data=form)
    assert response.status_code == 422 and message in response.text
    assert repo.list_templates(db_of(data_dir)) == []


def test_manual_template_custom_error_keeps_the_custom_choice_selected(client):
    form = {"name": "t", "header": ["Status", "Kept"], "field": ["__custom__", "__custom__"], "field_text": ["", "Typed"]}
    response = client.post("/templates/manual", data=form)
    assert response.status_code == 422
    assert response.text.count('<option value="__custom__" data-custom selected>') == 2
    assert 'value="Typed">' in response.text


def test_mapping_page_custom_text_round_trips(client, data_dir):
    template_id = int(upload_template(client).headers["location"].rsplit("/", 1)[1])
    headers = list(ODOO_SALE_ORDER_COLUMNS)
    form = {f"map_{index}": path for index, path in enumerate(ODOO_SALE_ORDER_COLUMNS.values())}
    form.update({"map_3": "__custom__", "map_text_3": " Sales Team ", "map_8": "", "map_text_8": "ignored"})
    assert client.post(f"/templates/{template_id}/mapping", data=form, follow_redirects=False).status_code == 303
    column_map = repo.get_template(db_of(data_dir), template_id).column_map
    assert column_map[headers[3]] == "literal:Sales Team" and headers[8] not in column_map

    page = client.get(f"/templates/{template_id}").text
    assert page.count('<option value="__custom__" data-custom selected>') == 1
    assert re.search(r'name="map_text_3"[^>]*value="Sales Team">', page)
    # Only the custom row's text input is visible.
    assert len(re.findall(r'name="map_text_\d+"[^>]*>', page)) == len(headers)
    assert len(re.findall(r'name="map_text_\d+"[^>]*hidden>', page)) == len(headers) - 1


def test_mapping_page_custom_without_text_is_rejected_and_kept(client, data_dir):
    template_id = int(upload_template(client).headers["location"].rsplit("/", 1)[1])
    saved_map = repo.get_template(db_of(data_dir), template_id).column_map
    response = client.post(f"/templates/{template_id}/mapping", data={"map_0": "__custom__", "map_text_0": " "})
    assert response.status_code == 422
    assert "Enter a value for the custom column &#34;Order Reference&#34;" in response.text
    assert '<option value="__custom__" data-custom selected>' in response.text
    assert repo.get_template(db_of(data_dir), template_id).column_map == saved_map


def test_export_without_item_rows_is_blocked_with_a_plain_reason(client, data_dir, tmp_path):
    conversion_id = ready_for_export(client, data_dir)
    for _ in range(2):
        assert client.delete(f"{order_url(conversion_id)}/lines/0", headers=HX).status_code == 200
    review = client.get(order_url(conversion_id)).text
    assert "No item rows were found" in review
    confirm_everything(client, conversion_id)

    preview = client.get(f"/conversions/{conversion_id}/preview")
    assert preview.status_code == 200 and 'name="acknowledge_unchecked"' in preview.text
    blocked = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx"})
    assert blocked.status_code == 422 and "spreadsheetml" not in blocked.headers["content-type"]
    assert "No item rows were found" in blocked.text and "have not been checked" not in blocked.text
    assert 'name="acknowledge_unchecked"' in blocked.text

    allowed = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx", "acknowledge_unchecked": "1"})
    assert allowed.status_code == 200 and allowed.headers["content-type"].startswith("application/vnd.openxmlformats")


def test_batch_export_combines_every_done_page_in_page_order(client, data_dir, three_pages, tmp_path):
    conversion_id = ready_for_export(client, data_dir, name="Davao week 1 #2.pdf")

    preview = client.get(f"/conversions/{conversion_id}/preview")
    assert preview.status_code == 200
    batch = preview.context["preview"]
    assert batch.order_count == 2 and batch.skipped_pages == [2] and batch.failed_pages == []
    assert [a.page_number for a in batch.orders_needing_attention] == [1, 3]
    assert batch.row_pages[0] == 1 and batch.row_pages[-1] == 3 and len(batch.row_pages) == len(batch.rows)

    response = client.post(f"/conversions/{conversion_id}/export", data={"format": "xlsx", "acknowledge_unchecked": "1"})

    assert response.status_code == 200
    assert 'filename="Davao_week_1__2.xlsx"' in response.headers["content-disposition"]
    output = tmp_path / "batch.xlsx"
    output.write_bytes(response.content)
    sheet = load_workbook(output)["Converted"]
    refs = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    # Order header fields repeat on every line row.
    assert list(dict.fromkeys(refs)) == ["S08824", "S07002"] and refs[0] == refs[1] == "S08824"
    assert sheet.max_row - 1 == len(batch.rows)


def test_skipped_pages_never_block_the_export(client, data_dir, three_pages):
    conversion_id = ready_for_export(client, data_dir)
    client.post(f"{order_url(conversion_id, 3)}/skip", headers=HX)
    confirm_everything(client, conversion_id, 1)

    preview = client.get(f"/conversions/{conversion_id}/preview").context["preview"]
    assert preview.skipped_pages == [2, 3] and not preview.needs_acknowledgement

    response = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv"})
    assert response.status_code == 200
    assert 'filename="S08824.csv"' in response.headers["content-disposition"]


def test_failed_page_blocks_export_until_acknowledged(client, data_dir, three_pages):
    three_pages["errors"] = {3: OcrToolError("tesseract exited 1")}
    conversion_id = ready_for_export(client, data_dir)
    confirm_everything(client, conversion_id, 1)

    blocked = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv"})

    assert blocked.status_code == 422
    # assert "Page 3 could not be read, so their orders are not in this download." in blocked.context["export_error"]
    assert "Page 3 could not be read, so its order is not in this download." in blocked.context["export_error"]
    assert blocked.context["failed_pages"] == [3] and blocked.context["needs_acknowledgement"]
    allowed = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv", "acknowledge_unchecked": "1"})
    assert allowed.status_code == 200


def test_page_requeued_after_the_batch_finished_blocks_export(client, data_dir, three_pages):
    conversion_id = ready_for_export(client, data_dir)
    db_path = db_of(data_dir)
    # The race QA found: a retry's requeue lands before the conversion status goes back to queued.
    repo.set_page_status(db_path, conversion_id, 3, "queued")
    assert repo.get_conversion(db_path, conversion_id).status == "done"
    expected = "Page 3 is still being read. Wait for it to finish before exporting, so its order is included."

    blocked = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv", "acknowledge_unchecked": "1"})

    assert blocked.status_code == 422 and not blocked.headers["content-type"].startswith("text/csv")
    assert blocked.context["export_error"] == expected
    assert not (repo.conversion_dir(data_dir.resolve(), conversion_id) / "exports").exists()
    preview = client.get(f"/conversions/{conversion_id}/preview")
    assert preview.status_code == 422 and preview.context["export_error"] == expected
    assert not client.get(f"/conversions/{conversion_id}").context["can_export"]

    repo.set_page_status(db_path, conversion_id, 2, "processing")
    plural = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv", "acknowledge_unchecked": "1"})
    assert plural.context["export_error"].startswith("Pages 2, 3 are still being read. Wait for them to finish")


def test_duplicate_order_ref_blocks_export_until_acknowledged(client, data_dir, ocr_pages):
    ocr_pages["pages"] = 2
    conversion_id = ready_for_export(client, data_dir)
    confirm_everything(client, conversion_id, 1)
    confirm_everything(client, conversion_id, 2)

    blocked = client.post(f"/conversions/{conversion_id}/export", data={"format": "csv"})

    assert blocked.status_code == 422
    assert blocked.context["duplicate_order_refs"] == {"S08824": [1, 2]}
    assert "Order S08824 was read from more than one page" in blocked.context["export_error"]


def test_template_preview_and_export_need_an_extracted_order(client, data_dir, ocr_pages):
    ocr_pages["words"] = {1: parse_tesseract_tsv(GARBAGE_TSV)}
    conversion_id = uploaded_conversion(client)
    assert not client.get(f"/conversions/{conversion_id}").context["can_export"]
    template_id = int(upload_template(client).headers["location"].rsplit("/", 1)[1])

    for path in ("template", "preview"):
        response = client.get(f"/conversions/{conversion_id}/{path}", follow_redirects=False)
        assert response.status_code == 303 and response.headers["location"] == f"/conversions/{conversion_id}"
    assert client.post(f"/conversions/{conversion_id}/template", data={"template_id": str(template_id)}).status_code == 409
    assert client.post(f"/conversions/{conversion_id}/export", data={"format": "csv"}).status_code == 409


def test_template_and_preview_wait_for_the_whole_batch(client, data_dir):
    db_path = db_of(data_dir)
    conversion = repo.create_conversion(db_path, "scan.pdf", ".pdf")
    repo.create_pages(db_path, conversion.id, 2)
    repo.set_status(db_path, conversion.id, "processing")
    order = extract_order(fixture_words("odoo_sale_order_S08824.tsv"), load_profile(runner.SALES_ORDER_PROFILE))
    repo.save_extraction(db_path, conversion.id, 1, order, runner.SALES_ORDER_PROFILE)

    response = client.get(f"/conversions/{conversion.id}/template", follow_redirects=False)
    assert response.status_code == 303 and response.headers["location"] == f"/conversions/{conversion.id}"
    assert client.post(f"/conversions/{conversion.id}/export", data={"format": "csv"}).status_code == 409
    # Reviewing finished pages while the rest of the batch is still being read is allowed.
    assert client.get(order_url(conversion.id, 1)).status_code == 200


def test_delete_template(client, data_dir):
    location = upload_template(client).headers["location"]
    template_id = int(location.rsplit("/", 1)[1])
    template = repo.get_template(db_of(data_dir), template_id)
    assert (data_dir.resolve() / template.file_path).is_file()

    assert client.delete(f"/templates/{template_id}").status_code == 200
    assert not (data_dir.resolve() / template.file_path).exists()


def test_startup_recovers_interrupted_conversions(data_dir, ocr_pages):
    db_path = db_of(data_dir)
    init_db(db_path)
    stuck = repo.create_conversion(db_path, "scan.pdf", ".pdf")
    repo.create_pages(db_path, stuck.id, 2)
    repo.set_status(db_path, stuck.id, "processing")
    repo.set_page_status(db_path, stuck.id, 1, "not_sales_order")
    repo.set_page_status(db_path, stuck.id, 2, "processing")

    with TestClient(create_app(data_dir, executor_factory=InlineExecutor)):
        recovered = repo.get_conversion(db_path, stuck.id)
        pages = repo.list_pages(db_path, stuck.id)

    assert (recovered.status, recovered.error) == ("failed", INTERRUPTED_ERROR)
    assert [(p.status, p.error) for p in pages] == [("not_sales_order", None), ("failed", PAGE_INTERRUPTED_ERROR)]


def test_page_whose_failure_cannot_be_recorded_is_logged_and_recovered_on_restart(
    data_dir, ocr_pages, monkeypatch, caplog
):
    ocr_pages["pages"] = 2
    ocr_pages["errors"] = {1: OcrToolError("tesseract exited 1")}
    db_path = db_of(data_dir)
    real_set_page_status = repo.set_page_status

    def locked_on_failure(db, conversion_id, page_number, status, error=None):
        if status == repo.FAILED:
            raise sqlite3.OperationalError("database is locked")
        real_set_page_status(db, conversion_id, page_number, status, error)

    monkeypatch.setattr(repo, "set_page_status", locked_on_failure)
    with caplog.at_level(logging.ERROR, logger="converter.worker.runner"), \
            TestClient(create_app(data_dir, executor_factory=InlineExecutor)) as test_client:
        conversion_id = uploaded_conversion(test_client)
    runner.reset_stop_request()

    assert "could not record the page as failed" in caplog.text
    # The rest of the batch still got read.
    assert [p.status for p in repo.list_pages(db_path, conversion_id)] == ["processing", "done"]
    assert repo.get_conversion(db_path, conversion_id).status == "done"

    monkeypatch.setattr(repo, "set_page_status", real_set_page_status)
    ocr_pages["errors"] = {}
    with TestClient(create_app(data_dir, executor_factory=InlineExecutor)) as test_client:
        recovered = page_of(data_dir, conversion_id, 1)
        assert (recovered.status, recovered.error) == ("failed", PAGE_INTERRUPTED_ERROR)
        assert test_client.post(f"{order_url(conversion_id, 1)}/retry", headers=HX).status_code == 200
        assert page_of(data_dir, conversion_id, 1).status == "done"
    runner.reset_stop_request()


def test_shutdown_requests_stop_and_startup_clears_it(data_dir, ocr_pages):
    with TestClient(create_app(data_dir, executor_factory=InlineExecutor)) as test_client:
        assert not runner._stop_requested.is_set()
        conversion_id = uploaded_conversion(test_client)
    assert runner._stop_requested.is_set()

    with TestClient(create_app(data_dir, executor_factory=InlineExecutor)):
        assert not runner._stop_requested.is_set()
    runner.reset_stop_request()
    assert repo.get_conversion(db_of(data_dir), conversion_id).status == "done"


def test_responses_carry_content_security_policy(client):
    assert "script-src 'self'" in client.get("/").headers["Content-Security-Policy"]


def test_route_handlers_are_sync():
    """Handlers touch SQLite/filesystem/openpyxl; an async def handler would block the event loop."""
    for module in ("routes_conversions.py", "routes_templates.py"):
        assert "async def" not in (APP_DIR / module).read_text(encoding="utf-8")
