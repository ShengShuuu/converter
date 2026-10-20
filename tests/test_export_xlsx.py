from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart

from converter.export.xlsx import write_xlsx
from converter.mapping import ODOO_SALE_ORDER_COLUMNS
from converter.models import Order, OrderLine

TEMPLATE_HEADERS = [
    "Order Reference",
    "Creation Date",
    "Customer",
    "Salesperson",
    "Payment Terms",
    "Order Lines/Product",
    "Order Lines/Quantity",
    "Order Lines/Subtotal",
    "Order Lines/Taxes",
]


def read_rows(path, sheet="Converted"):
    return [list(row) for row in load_workbook(path)[sheet].iter_rows(values_only=True)]


def test_synthetic_extraction_exports_one_row_per_line(synthetic_order, template_path, tmp_path):
    output = write_xlsx([synthetic_order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    header, *rows = read_rows(output)
    order_date = datetime(2026, 4, 22, 10, 41, 30)
    assert header == TEMPLATE_HEADERS
    # Replaced: customer and payment_terms are now extracted, no longer blank.
    # assert rows == [
    #     ["S08824", order_date, None, "MARK LUISEN", None,
    #      "FSHB100 / 100W F-SERIES CONVENTIONAL HIGHBAY 24 MONTHS WARRANTY", 1, 2600, None],
    #     ["S08824", order_date, None, "MARK LUISEN", None,
    #      "PX-S100 / 100W PX-S SOLAR FLOOD LIGHT", 2, 5760, "Output Tax (12%)"],
    # ]
    header_fields = ["S08824", order_date, "ACME TRADING CORP.", "MARK LUISEN", "POS / THERMAL PAYMENT"]
    assert rows == [
        [*header_fields, "FSHB100 / 100W F-SERIES CONVENTIONAL HIGHBAY 24 MONTHS WARRANTY", 1, 2600, None],
        [*header_fields, "PX-S100 / 100W PX-S SOLAR FLOOD LIGHT", 2, 5760, "Output Tax (12%)"],
    ]


def test_cells_are_typed_not_strings(synthetic_order, template_path, tmp_path):
    output = write_xlsx([synthetic_order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    worksheet = load_workbook(output)["Converted"]
    assert worksheet["B2"].is_date
    assert worksheet["B2"].number_format == "mmmm d, yyyy"
    assert worksheet["G2"].data_type == "n"
    assert worksheet["H2"].data_type == "n"


def test_shape_matches_real_sample_rows(template_path, tmp_path):
    """Mirrors S07001 from the user's real output: header repeated per line, blank subtotal/taxes kept blank."""
    order = Order(
        order_ref="S07001",
        order_date=datetime(2025, 11, 26),
        customer="GIAN ALEXIS BORGONIA",
        salesperson="JOACHAS SANTOS",
        payment_terms="GCASH PAYMENT",
        lines=[
            OrderLine(product="PX-S100", quantity=Decimal("2"), subtotal=Decimal("11260")),
            OrderLine(product="ACFL100 EC", quantity=Decimal("4")),
            OrderLine(product="COBRA50", quantity=Decimal("1")),
        ],
    )
    header_fields = ["S07001", datetime(2025, 11, 26), "GIAN ALEXIS BORGONIA", "JOACHAS SANTOS", "GCASH PAYMENT"]
    output = write_xlsx([order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    _, *rows = read_rows(output)
    assert rows == [
        [*header_fields, "PX-S100", 2, 11260, None],
        [*header_fields, "ACFL100 EC", 4, None, None],
        [*header_fields, "COBRA50", 1, None, None],
    ]


def test_template_sample_rows_replaced_and_formatting_kept(synthetic_order, template_path, tmp_path):
    template_widths = {k: v.width for k, v in load_workbook(template_path)["Converted"].column_dimensions.items()}
    output = write_xlsx([synthetic_order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    worksheet = load_workbook(output)["Converted"]
    assert worksheet.max_row == 3
    assert "S06977" not in [cell.value for cell in worksheet["A"]]
    assert {k: v.width for k, v in worksheet.column_dimensions.items()} == template_widths


def test_leftover_template_sheets_are_not_exported(synthetic_order, template_path, tmp_path):
    assert len(load_workbook(template_path).sheetnames) > 1
    output = write_xlsx([synthetic_order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    assert load_workbook(output).sheetnames == ["Converted"]


def test_only_header_sheet_exported_when_it_is_not_first(synthetic_order, tmp_path):
    template = Workbook()
    stale = template.active
    stale.title = "Old Export"
    stale.append(["S06977", "stale row without header"])
    target = template.create_sheet("Orders")
    target.append(TEMPLATE_HEADERS)
    template.create_chartsheet("Chart").add_chart(BarChart())
    template.save(tmp_path / "template.xlsx")

    output = write_xlsx([synthetic_order], tmp_path / "template.xlsx", tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    exported = load_workbook(output)
    assert exported.sheetnames == ["Orders"]
    assert exported.active.title == "Orders"
    assert len(read_rows(output, sheet="Orders")) == 3


def test_order_without_lines_still_exported(template_path, tmp_path):
    output = write_xlsx([Order(order_ref="S1")], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    _, *rows = read_rows(output)
    assert rows == [["S1", None, None, None, None, None, None, None, None]]


def test_formula_like_text_is_written_as_plain_text(template_path, tmp_path):
    product = '=HYPERLINK("http://evil.example")'
    order = Order(order_ref="=1+1", lines=[OrderLine(product=product)])
    output = write_xlsx([order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    worksheet = load_workbook(output)["Converted"]
    assert (worksheet["A2"].value, worksheet["A2"].data_type) == ("=1+1", "s")
    assert (worksheet["F2"].value, worksheet["F2"].data_type) == (product, "s")


def test_control_characters_are_stripped_instead_of_failing(template_path, tmp_path):
    order = Order(customer="ACME\x0bTRADING", lines=[OrderLine(product="LAMP\x00 100W")])
    output = write_xlsx([order], template_path, tmp_path / "out.xlsx", ODOO_SALE_ORDER_COLUMNS)
    worksheet = load_workbook(output)["Converted"]
    assert (worksheet["C2"].value, worksheet["F2"].value) == ("ACMETRADING", "LAMP 100W")
